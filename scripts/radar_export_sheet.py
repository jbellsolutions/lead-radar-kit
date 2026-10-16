"""Export one run's leads to a new, private Google Sheet (a minimal slice of M4's Sheets exporter).

Tabs:
- Leads: buyers only (sellers dropped, per outputs.drop_sellers), High -> Medium -> Low, newest first.
- By platform: fetched / classified / tiers / sellers / rule drops / errors per source, with two charts.
- Run: window, credits, model cost and caveats.

Uses the `gws` CLI (your own Google login). Never shares the sheet; it stays private to the account that creates it.

    PYTHONPATH=src .venv/bin/python scripts/radar_export_sheet.py <radar> [--run <run_id>|latest] [--title ...]
"""
from __future__ import annotations

import argparse
import json
import subprocess
import sys
from collections import Counter, defaultdict

from lead_radar import config
from lead_radar.db import Store

TIER_ORDER = {"high": 0, "medium": 1, "low": 2}
LEAD_HEADER = ["Platform", "Tier", "Urgency", "Posted (UTC)", "Author", "Post URL", "Evidence (verbatim)",
               "Need", "Decided by", "Confidence", "Post text (first 300 chars)",
               "Contact", "Name", "Company", "Email", "Email check", "Phone", "LinkedIn"]
CONTACT_LABEL = {"verified": "Verified", "accept_all": "Accept-all", "not_found": "No email found", "pending": "Pending",
                 "unresolvable": "Unresolvable", "failed": "Lookup failed", "none": ""}
DEFAULT_CAVEATS = [
    "Classifier accuracy is only measured on the virtual-assistant offer; spot-read every lead before acting.",
    "Job posts (a business hiring for the role) can appear as leads; that boundary is still under review.",
    "X is covered only via Google site:x.com; Skool and private Facebook groups need the browser lane (not built).",
    "Public posts only; nobody was contacted.",
]
BASES_NOTE = ("Fetched, Credits and Errors count by search source; Classified, tiers, Sellers and Rule drops count by "
              "each post's own platform (Google results are counted under the site they came from).")
PLATFORM_HEADER = ["Platform", "Fetched", "Classified", "High", "Medium", "Low", "Leads (H+M+L)",
                   "Sellers (dropped)", "Rule drops", "Credits", "Errors"]


def lead_rows(rows: list[dict], contacts: dict | None = None) -> list[list]:
    """rows: joined classification + post dicts. Buyers with a tier only, ordered by tier then recency.
    contacts: subject -> contact card (Store.get_contact); the contact columns stay empty without it."""
    from lead_radar.enrich import badge, subject_of
    leads = [r for r in rows if r["role"] == "buyer" and r["tier"] in TIER_ORDER]
    leads.sort(key=lambda r: (TIER_ORDER[r["tier"]], -(_ts(r.get("posted_at")))))
    out = []
    for r in leads:
        c = (contacts or {}).get(subject_of(r["platform"], "post", r.get("author_handle"), r.get("author_url"))) or {}
        out.append([r["platform"], r["tier"].title(), r["urgency"], (r.get("posted_at") or "")[:16].replace("T", " "),
                    r.get("author_handle") or r.get("author_name") or "", r["url"], r["evidence_quote"], r["need_summary"],
                    r["decided_by"], round(float(r["confidence"]), 2), " ".join((r.get("text") or "").split())[:300],
                    CONTACT_LABEL[badge(c or None)], c.get("name") or "", c.get("company") or "", c.get("email") or "",
                    c.get("email_verdict") or "", c.get("phone") or "", c.get("linkedin_url") or ""])
    return out


def platform_rows(rows: list[dict], drops: list[dict], per_source: dict) -> list[list]:
    tiers: dict[str, Counter] = defaultdict(Counter)
    for r in rows:
        tiers[r["platform"]]["classified"] += 1
        if r["role"] == "seller":
            tiers[r["platform"]]["seller"] += 1
        elif r["role"] == "buyer" and r["tier"] in TIER_ORDER:
            tiers[r["platform"]][r["tier"]] += 1
    dropped = Counter(d["platform"] for d in drops)
    # Google results are re-attributed to their target site, so source (fetch) and platform (post) differ.
    platforms = sorted(set(per_source) | set(tiers) | set(dropped))
    out = []
    for p in platforms:
        t, src = tiers[p], per_source.get(p) or {}
        leads = t["high"] + t["medium"] + t["low"]
        out.append([p, src.get("posts", 0), t["classified"], t["high"], t["medium"], t["low"], leads, t["seller"],
                    dropped[p], src.get("credits", 0), src.get("errors", 0)])
    return out


def _ts(v: str | None) -> float:
    from datetime import datetime
    try:
        return datetime.fromisoformat(v).timestamp() if v else 0.0
    except ValueError:
        return 0.0


def gws(*args: str, body: dict | None = None) -> dict:
    cmd = ["gws", "sheets", "spreadsheets", *args] + (["--json", json.dumps(body)] if body is not None else [])
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode:
        raise SystemExit(f"gws failed: {res.stderr.strip()[:500]}")
    return json.loads(res.stdout) if res.stdout.strip() else {}


def column_chart(sheet_id: int, title: str, n_rows: int, series_cols: list[int], stacked: bool, anchor_row: int) -> dict:
    def src(col):
        return {"sourceRange": {"sources": [{"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": n_rows + 1,
                                             "startColumnIndex": col, "endColumnIndex": col + 1}]}}
    return {"addChart": {"chart": {
        "spec": {"title": title, "basicChart": {
            "chartType": "COLUMN", "legendPosition": "BOTTOM_LEGEND", "headerCount": 1,
            "stackedType": "STACKED" if stacked else "NOT_STACKED",
            "axis": [{"position": "BOTTOM_AXIS", "title": "Platform"}, {"position": "LEFT_AXIS", "title": "Posts"}],
            "domains": [{"domain": src(0)}],
            "series": [{"series": src(c), "targetAxis": "LEFT_AXIS"} for c in series_cols]}},
        "position": {"overlayPosition": {"anchorCell": {"sheetId": sheet_id, "rowIndex": anchor_row, "columnIndex": 0},
                                         "widthPixels": 760, "heightPixels": 380}}}}}


def _safe(row: list) -> list:
    """Text that starts like a formula would execute in a spreadsheet; post text is untrusted."""
    return [("'" + v) if isinstance(v, str) and v and v[0] in "=+-@" else v for v in row]


def write_csv(path: str, title: str, leads: list[list], plat: list[list], run_tab: list[list]) -> None:
    """One tab for a Drive upload: per-platform table with in-cell SPARKLINE bar charts, then leads, then run notes."""
    import csv
    first = 5  # row of the first platform (title, blank, section label, header come first)
    mx = max([r[6] for r in plat] + [1])
    mx2 = max([r[6] + r[7] for r in plat] + [1])
    with open(path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow([title])
        w.writerow([])
        w.writerow(["BY PLATFORM"])
        w.writerow(PLATFORM_HEADER + ["Leads chart: High (green) / Medium (amber) / Low (grey)",
                                      "Sellers (red) vs leads (green)"])
        for i, r in enumerate(plat, start=first):
            w.writerow(r + [
                f'=SPARKLINE({{D{i},E{i},F{i}}},{{"charttype","bar";"max",{mx};"color1","#1a7f37";"color2","#f0a500";'
                f'"color3","#9aa0a6"}})',
                f'=SPARKLINE({{H{i},G{i}}},{{"charttype","bar";"max",{mx2};"color1","#c5221f";"color2","#1a7f37"}})'])
        end = first + len(plat) - 1
        w.writerow(["TOTAL"] + [f"=SUM({c}{first}:{c}{end})" for c in "BCDEFGHIJK"])
        w.writerow([BASES_NOTE])
        w.writerow([])
        w.writerow(["LEADS (buyers only; sellers dropped)"])
        w.writerow(LEAD_HEADER)
        w.writerows(_safe(r) for r in leads)
        w.writerow([])
        w.writerow(["RUN"])
        w.writerows(_safe(r) for r in run_tab)


def write_xlsx(path: str, leads: list[list], plat: list[list], run_tab: list[list], hours: int) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    for row in [LEAD_HEADER] + leads:
        ws.append(row)
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHIJK", (10, 8, 9, 16, 22, 48, 60, 50, 10, 10, 80)):
        ws.column_dimensions[col].width = width
    wp = wb.create_sheet("By platform")
    for row in [PLATFORM_HEADER] + plat:
        wp.append(row)
    wp.freeze_panes = "A2"
    for s in (ws, wp):
        for c in s[1]:
            c.font = Font(bold=True)
    n = len(plat) + 1
    cats = Reference(wp, min_col=1, min_row=2, max_row=n)

    def chart(title, cols, stacked, anchor):
        ch = BarChart()
        ch.type, ch.title, ch.width, ch.height = "col", title, 22, 10
        ch.y_axis.title, ch.x_axis.title = "Posts", "Platform"
        if stacked:
            ch.grouping, ch.overlap = "stacked", 100
        for c in cols:
            ch.add_data(Reference(wp, min_col=c, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(cats)
        wp.add_chart(ch, anchor)
    chart(f"Leads by platform and tier (last {hours}h)", [4, 5, 6], True, f"A{n + 3}")
    chart("Posts classified vs leads vs sellers, by platform", [3, 7, 8], False, f"A{n + 25}")
    wr = wb.create_sheet("Run")
    for row in run_tab:
        wr.append(row)
    wr.column_dimensions["A"].width, wr.column_dimensions["B"].width = 60, 90
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("radar")
    ap.add_argument("--run", default="latest",
                    help="run id, or several comma-separated to combine passes (default: the radar's latest run)")
    ap.add_argument("--title")
    ap.add_argument("--caveat", action="append", default=[], help="extra caveat (the defaults are always included)")
    ap.add_argument("--xlsx", help="write an .xlsx (native charts) here instead of calling gws")
    ap.add_argument("--csv", help="write a one-tab CSV (summary with SPARKLINE charts, leads, run notes) here instead of "
                                  "calling gws; upload it through a Google Drive connector with conversion to a Sheet")
    a = ap.parse_args()

    settings = config.load()
    store = Store.open(settings.db_path)
    radar = store.get_radar(a.radar)
    if not radar:
        raise SystemExit("no such radar")
    q = store.conn.execute
    if a.run == "latest":  # the radar's latest run, not the store's
        row = q("SELECT id FROM runs WHERE radar_id=? ORDER BY started_at DESC LIMIT 1", (radar.id,)).fetchone()
        run_ids = [row["id"]] if row else []
    else:  # comma-separated to combine several passes over the same window
        run_ids = [x.strip() for x in a.run.split(",") if x.strip()]
    runs = [store.get_run(r) for r in run_ids]
    if not runs or any(r is None or r.radar_id != radar.id for r in runs):
        raise SystemExit("no such run for this radar")
    marks = ",".join("?" * len(run_ids))
    rows = [dict(r) for r in q(
        "SELECT DISTINCT c.*, p.platform, p.url, p.author_handle, p.author_name, p.author_url, p.posted_at, p.text "
        f"FROM classifications c JOIN run_posts rp ON rp.post_id = c.post_id AND rp.run_id IN ({marks}) "
        "JOIN posts p ON p.id = c.post_id WHERE c.radar_id = ?", (*run_ids, radar.id))]
    drops = [dict(r) for r in q("SELECT d.reason, p.platform FROM drops d JOIN posts p ON p.id = d.post_id "
                                f"WHERE d.run_id IN ({marks}) AND d.radar_id = ?", (*run_ids, radar.id))]
    per_source: dict = {}
    for r in runs:
        for src, v in r.per_source.items():
            agg = per_source.setdefault(src, Counter())
            agg.update({k: v.get(k, 0) for k in ("posts", "credits", "errors")})
    brief = store.get_brief(radar.intent_brief_id)
    leads = lead_rows(rows, store.contacts_for([radar.id]))
    plat = platform_rows(rows, drops, per_source)
    first = min(r.started_at for r in runs)

    title = a.title or f"Lead Radar - {radar.name} - {first:%Y-%m-%d} ({radar.lookback_days * 24}h)"
    credits = sum(r.credits_used for r in runs)
    credits_usd = credits * float(settings.get("scrapecreators.credit_usd", 0) or 0)
    total_usd = sum(r.cost_usd for r in runs)

    def total(k):
        return sum((r.counts.get(k) or 0) for r in runs)
    tiers = Counter(x[1] for x in leads)
    run_tab = [["Radar", radar.name], ["Offer", brief.offer], ["Runs (passes)", ", ".join(run_ids)],
               ["Started (UTC)", f"{first:%Y-%m-%d %H:%M}"], ["Lookback", f"{radar.lookback_days * 24} hours"],
               ["Posts fetched", total("fetched")], ["New unique posts", total("to_classify")],
               ["Dropped by rules (older than window, too short, creator's own post, templates)", total("stage0_dropped")],
               ["Classified", len(rows)], ["Leads (buyers, H+M+L)", len(leads)],
               ["High", tiers["High"]], ["Medium", tiers["Medium"]], ["Low", tiers["Low"]],
               ["Sellers (dropped from Leads)", sum(1 for r in rows if r["role"] == "seller")],
               ["Decided by Jev", f"{sum(1 for r in rows if r['decided_by'] == 'jev')}/{len(rows)}"],
               ["ScrapeCreators credits", credits],
               ["ScrapeCreators cost (USD)", round(credits_usd, 4)],
               ["Model cost (USD)", round(total_usd - credits_usd, 4)],
               ["Total cost (USD)", round(total_usd, 4)],
               ["Note", BASES_NOTE]]
    run_tab += [["Caveat", x] for x in DEFAULT_CAVEATS + a.caveat]
    if a.xlsx or a.csv:
        if a.xlsx:
            write_xlsx(a.xlsx, leads, plat, run_tab, radar.lookback_days * 24)
        if a.csv:
            write_csv(a.csv, title, leads, plat, run_tab)
        print(json.dumps({"xlsx": a.xlsx, "csv": a.csv, "title": title, "leads": len(leads), "platforms": plat}))
        return
    created = gws("create", body={"properties": {"title": title}, "sheets": [
        {"properties": {"title": "Leads", "gridProperties": {"frozenRowCount": 1}}},
        {"properties": {"title": "By platform", "gridProperties": {"frozenRowCount": 1}}},
        {"properties": {"title": "Run"}}]})
    sid = created["spreadsheetId"]
    ids = {s["properties"]["title"]: s["properties"]["sheetId"] for s in created["sheets"]}

    gws("values", "batchUpdate", "--params", json.dumps({"spreadsheetId": sid}), body={
        "valueInputOption": "RAW", "data": [
            {"range": "Leads!A1", "values": [LEAD_HEADER] + leads},
            {"range": "'By platform'!A1", "values": [PLATFORM_HEADER] + plat},
            {"range": "Run!A1", "values": run_tab}]})

    n = len(plat)
    gws("batchUpdate", "--params", json.dumps({"spreadsheetId": sid}), body={"requests": [
        column_chart(ids["By platform"], "Leads by platform and tier (last 72h)", n, [3, 4, 5], True, n + 3),
        column_chart(ids["By platform"], "Posts classified vs leads vs sellers, by platform", n, [2, 6, 7], False, n + 24),
        {"repeatCell": {"range": {"sheetId": ids["Leads"], "startRowIndex": 0, "endRowIndex": 1},
                        "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}}, "fields": "userEnteredFormat.textFormat.bold"}},
        {"repeatCell": {"range": {"sheetId": ids["By platform"], "startRowIndex": 0, "endRowIndex": 1},
                        "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}}, "fields": "userEnteredFormat.textFormat.bold"}},
        {"autoResizeDimensions": {"dimensions": {"sheetId": ids["By platform"], "dimension": "COLUMNS", "startIndex": 0, "endIndex": 11}}},
        {"updateDimensionProperties": {"range": {"sheetId": ids["Leads"], "dimension": "COLUMNS", "startIndex": 6, "endIndex": 8},
                                       "properties": {"pixelSize": 360}, "fields": "pixelSize"}},
    ]})
    print(json.dumps({"spreadsheet_id": sid, "url": f"https://docs.google.com/spreadsheets/d/{sid}/edit",
                      "leads": len(leads), "platforms": plat}))


if __name__ == "__main__":
    sys.exit(main())
